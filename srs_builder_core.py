# -*- coding: utf-8 -*-
"""
Generator Skrip Dokumen Software Requirements Specification (SRS)
Standar IEEE 830-1998 untuk:
SISTEM PENDUKUNG KEPUTUSAN PEMILIHAN LAPTOP GAMING UNTUK MAHASISWA
MENGGUNAKAN METODE TOPSIS BERBASIS WEB
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from srs_helpers import (
    set_cell_shading,
    set_cell_margins,
    set_table_borders,
    add_page_number_to_run,
    add_total_pages_to_run,
)

# Palet Warna Dokumen Eksekutif
NAVY_PRIMARY = RGBColor(15, 44, 89)      # #0F2C59 - Judul & Header Utama
BLUE_SECONDARY = RGBColor(51, 124, 207)  # #337CCF - Sub-heading & Aksen
SLATE_DARK = RGBColor(30, 41, 59)        # #1E293B - Teks Tubuh Dokumen
GRAY_MUTED = RGBColor(100, 116, 139)     # #64748B - Teks Sekunder/Caption
HEX_PRIMARY = "0F2C59"
HEX_SECONDARY = "337CCF"
HEX_LIGHT_BG = "F8FAFC"
HEX_CALLOUT_BG = "F0F4F8"
HEX_BORDER = "CBD5E1"

def set_run_font(run, font_name="Calibri", size_pt=10.5, color=SLATE_DARK, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic

def add_p(doc, text="", bold_prefix=None, space_after=6, space_before=0, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, bold=True, color=NAVY_PRIMARY)
    
    if text:
        r_text = p.add_run(text)
        set_run_font(r_text)
    return p

def add_bullet(doc, text, bold_prefix=None, space_after=3):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, bold=True, color=NAVY_PRIMARY)
    if text:
        r_text = p.add_run(text)
        set_run_font(r_text)
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, font_name="Calibri", size_pt=15, color=NAVY_PRIMARY, bold=True)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, font_name="Calibri", size_pt=12.5, color=BLUE_SECONDARY, bold=True)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, font_name="Calibri", size_pt=11, color=NAVY_PRIMARY, bold=True)
    return p

def add_callout(doc, text, title="CATATAN TEKNIS PENTING", bg_color=HEX_CALLOUT_BG, border_color=HEX_PRIMARY):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_color)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"📌 {title}: ")
    set_run_font(r_title, font_name="Calibri", size_pt=10, color=NAVY_PRIMARY, bold=True)
    
    r_text = p.add_run(text)
    set_run_font(r_text, font_name="Calibri", size_pt=9.5, color=SLATE_DARK)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_formula_box(doc, formula_str, explanation=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F1F5F9")
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_SECONDARY}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=140)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_f = p.add_run(formula_str)
    set_run_font(r_f, font_name="Cambria Math", size_pt=11.5, color=NAVY_PRIMARY, bold=True)
    
    if explanation:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        r_exp = p2.add_run(f"Keterangan: {explanation}")
        set_run_font(r_exp, font_name="Calibri", size_pt=9, color=RGBColor(71, 85, 105), italic=True)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def create_styled_table(doc, headers, rows_data, col_widths=None, header_bg=HEX_PRIMARY, alt_bg=HEX_LIGHT_BG):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Set borders
    set_table_borders(table, color=HEX_BORDER, sz="4", val="single")
    
    # Format Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for c_idx, title in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        set_cell_shading(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_run_font(r, font_name="Calibri", size_pt=9.5, color=RGBColor(255, 255, 255), bold=True)
        
    # Format Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        trPr_r = row._tr.get_or_add_trPr()
        trPr_r.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_col = alt_bg if (r_idx % 2 == 1) else "FFFFFF"
        
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg_col)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            set_run_font(r, font_name="Calibri", size_pt=9, color=SLATE_DARK)

    # Set column widths
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

print("Styling engine initialized.")
